from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from failure_paths.eventgraph.models import PathTable
from openpyxl import Workbook

from failure_paths import ExcelEventGraph


def create_example_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "scenario1"

    # Metadata section
    ws.append(["metadata"])
    ws.append(
        [
            "TRAJECT_ID",
            "M_VAN",
            "M_TOT",
            "Ondergrondscenario",
            "ScenarioKans",
            "LENGTE_VAK",
            "HR_locatie",
            "dijkvaknummer",
            "Vaknaam",
            "TYPE_WATERKERING",
        ]
    )
    ws.append(["16-X", 0.0, 100.0, "ScenarioA", 0.1, 1.23, "HR01", 42, "Vak 42", "type"])

    # Path table
    ws.append(["faalpadschema"])
    ws.append(
        [
            "Faalpad_ID",
            "Overslag",
            "Indirect_mechanisme",
            "Initiele_gebeurtenis",
            "Vervolggebeurtenis_1",
            "Vervolggebeurtenis_2",
            "Vervolggebeurtenis_3",
            "Vervolggebeurtenis_4",
            "Vervolggebeurtenis_5",
            "Vervolggebeurtenis_6",
            "Vervolggebeurtenis_7",
        ]
    )
    ws.append([1, "ja", "nee", "init", "event1", "event2", None, None, None, None, None])

    # Event table
    ws.append(["faalpadkansen"])
    ws.append(["Faalpad_ID", "Knoop_ID", "h", "Pf_h", "Beta_h"])
    event_rows = [
        (1, 1, 0.0, 0.01, 2.3263),
        (1, 1, 1.0, 0.015, 2.1701),
        (1, 2, 0.0, 0.02, 2.0537),
        (1, 2, 1.0, 0.025, 1.9599),
        (1, 3, 0.0, 0.03, 1.8808),
        (1, 3, 1.0, 0.035, 1.8125),
    ]
    for row in event_rows:
        ws.append(list(row))

    for sheet_name in ("FP_overslag", "FP_graverij"):
        freq_sheet = wb.create_sheet(sheet_name)
        freq_sheet.append(["h", "Pf_h"])
        freq_sheet.append([0.0, 0.5])
        freq_sheet.append([1.0, 0.6])

    wb.save(path)


def test_excel_eventgraph_loads_example(tmp_path: Path) -> None:
    excel_path = tmp_path / "scenario.xlsx"
    create_example_workbook(excel_path)

    eeg = ExcelEventGraph.load(excel_path, "scenario1")

    assert eeg.metadata.df.iloc[0]["ScenarioKans"] == 0.1
    assert (1, 1) in eeg.graph.nodes
    pf = eeg.graph_events.get_event_prob((1, 2), 0.0)
    assert pf == pytest.approx(0.02, rel=1e-2)


def test_excel_eventgraph_requires_keywords(tmp_path: Path) -> None:
    excel_path = tmp_path / "invalid.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "scenario1"
    ws.append(["faalpadschema"])
    wb.save(excel_path)

    with pytest.raises(ValueError):
        ExcelEventGraph.load(excel_path, "scenario1")


def test_normalize_paths_table_handles_indirect_mechanism() -> None:
    df_paths = pd.DataFrame(
        {
            "Faalpad_ID": [1],
            "Overslag": ["JA"],
            "Indirect_mechanisme": ["  Nee  "],
            "Initiele_gebeurtenis": ["init"],
            "Vervolggebeurtenis_1": ["e1"],
            "Vervolggebeurtenis_2": [pd.NA],
            "Vervolggebeurtenis_3": [pd.NA],
            "Vervolggebeurtenis_4": [pd.NA],
            "Vervolggebeurtenis_5": [pd.NA],
            "Vervolggebeurtenis_6": [pd.NA],
            "Vervolggebeurtenis_7": [pd.NA],
        }
    )
    paths = PathTable.from_dataframe(df_paths)
    normalized = ExcelEventGraph._normalize_paths_table(paths)
    assert normalized.indirect_column == "geen indirect mechanisme"
    assert normalized.overslag_values == ["ja"]
    assert normalized.df.index.name == "Faalpad_ID"
    assert normalized.group_columns == ["Overslag", "geen indirect mechanisme"]


def test_load_frequency_tables_detects_dynamic_sheets(tmp_path: Path) -> None:
    excel_path = tmp_path / "freqs.xlsx"
    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        pd.DataFrame({"dummy": [1]}).to_excel(writer, sheet_name="scenario1", index=False)
        freq_df = pd.DataFrame({"h": [0.0, 1.0], "Pf_h": [0.5, 0.6]})
        freq_df.to_excel(writer, sheet_name="FP_OVERSLAG", index=False)
        freq_df.to_excel(writer, sheet_name="FP_custom", index=False)

    with pd.ExcelFile(excel_path) as xlsx:
        freq_tables = ExcelEventGraph._load_frequency_tables(xlsx)

    assert sorted(freq_tables.keys()) == ["custom", "overslag"]
    for table in freq_tables.values():
        pd.testing.assert_frame_equal(
            table.df.reset_index(drop=True), pd.DataFrame({"h": [0.0, 1.0], "Pf_h": [0.5, 0.6]})
        )


def test_excel_eventgraph_errors_on_missing_frequency_table(tmp_path: Path) -> None:
    excel_path = tmp_path / "missing_freq.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "scenario1"

    ws.append(["metadata"])
    ws.append(
        [
            "TRAJECT_ID",
            "M_VAN",
            "M_TOT",
            "Ondergrondscenario",
            "ScenarioKans",
            "LENGTE_VAK",
            "HR_locatie",
            "dijkvaknummer",
            "Vaknaam",
            "TYPE_WATERKERING",
        ]
    )
    ws.append(["16-X", 0.0, 100.0, "Scenario", 0.1, 1.0, "HR", 1, "Vak", "type"])

    ws.append(["faalpadschema"])
    ws.append(
        [
            "Faalpad_ID",
            "Overslag",
            "Indirect_mechanisme",
            "Initiele_gebeurtenis",
            "Vervolggebeurtenis_1",
            "Vervolggebeurtenis_2",
            "Vervolggebeurtenis_3",
            "Vervolggebeurtenis_4",
            "Vervolggebeurtenis_5",
            "Vervolggebeurtenis_6",
            "Vervolggebeurtenis_7",
        ]
    )
    ws.append([1, "ja", "nee", "init", "event1", None, None, None, None, None, None])
    ws.append([2, "nee", "nee", "init", "event1", None, None, None, None, None, None])

    ws.append(["faalpadkansen"])
    ws.append(["Faalpad_ID", "Knoop_ID", "h", "Pf_h", "Beta_h"])
    ws.append([1, 1, 0.0, 0.01, 2.3263])
    ws.append([2, 1, 0.0, 0.01, 2.3263])
    ws.append([1, 2, 0.0, 0.02, 2.0])
    ws.append([2, 2, 0.0, 0.02, 2.0])

    wb.save(excel_path)

    with pytest.raises(ValueError, match="Frequency table.*overslag"):
        ExcelEventGraph.load(excel_path, "scenario1")
